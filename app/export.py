"""One edition's readers as an Excel workbook, for the admin panel's download.

Two sheets:

  Signed up   everyone whose payment for this edition cleared — who joined
              that month, what they paid, and the notes they left
  To post     everyone this edition's envelope goes to, which also takes in
              readers on longer plans who signed up in an earlier month

Everything typed by a reader goes in as plain text. openpyxl stores any string
starting with "=" as a formula, so a note written as "=HYPERLINK(...)" would
otherwise become a live formula in the sheet — that is forced back to text.
"""

from __future__ import annotations

import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import cycles, plans
from .db import fetch_all

# (header, key, width). `key` is a column on the row, or a callable.
def _amount(row: dict) -> str:
    if row.get("amount_minor") and row.get("currency"):
        return plans.display(row["amount_minor"], row["currency"])
    return ""


def _yes(value) -> str:
    return "Yes" if value else ""


ADDRESS_COLUMNS = [
    ("Address line 1", "address_line1", 32),
    ("Address line 2", "address_line2", 26),
    ("Landmark", "landmark", 22),
    ("City", "city", 16),
    ("State", "state", 16),
    ("PIN / postal code", "pincode", 12),
    ("Country", "country", 14),
]

PERSON_COLUMNS = [
    ("Reference", "reference", 13),
    ("First name", "first_name", 14),
    ("Last name", "last_name", 14),
    ("Email", "email", 28),
    ("Phone", "phone", 17),
    ("Instagram", "instagram", 18),
    ("Birthday", "birthdate", 12),
]

NOTE_COLUMNS = [
    ("Reader's note (what they'd love a letter about)", "interests_note", 40),
    ("Gift?", lambda r: _yes(r.get("is_gift")), 7),
    ("Gift message (for the card)", "gift_message", 36),
    ("Admin note", "admin_note", 24),
]

SIGNED_UP = [
    ("Paid on", "paid_on", 18),
    *PERSON_COLUMNS,
    *ADDRESS_COLUMNS,
    ("Plan (months)", "bought_months", 8),
    ("Amount paid", _amount, 11),
    ("Code used", "bought_code", 12),
    *NOTE_COLUMNS,
    ("Envelopes still owed", "deliveries_remaining", 9),
    ("Status", "status", 10),
    ("Razorpay payment ID", "razorpay_payment_id", 22),
]

TO_POST = [
    *PERSON_COLUMNS,
    *ADDRESS_COLUMNS,
    *NOTE_COLUMNS,
    ("Plan (months)", "plan_months", 8),
    ("Joined with", lambda r: cycles.month_name(r["cycle"]), 15),
    ("Envelopes still owed", "deliveries_remaining", 9),
]

HEADER_FILL = PatternFill("solid", fgColor="DDE3C6")


def _cell_value(value):
    if isinstance(value, datetime):
        # Excel has no time zones. Shown in Indian time, which is what the
        # rest of the club runs on.
        return value.astimezone(cycles.IST).replace(tzinfo=None)
    if isinstance(value, list):
        return " | ".join(str(v) for v in value)
    return value


def _sheet(wb: Workbook, title: str, columns: list, rows: list[dict], first: bool) -> None:
    ws = wb.active if first else wb.create_sheet()
    ws.title = title[:31]

    for col, (header, _key, width) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(col)].width = width

    for r, row in enumerate(rows, start=2):
        for col, (_header, key, _width) in enumerate(columns, start=1):
            value = _cell_value(key(row) if callable(key) else row.get(key))
            cell = ws.cell(row=r, column=col, value=value)
            if isinstance(value, str) and value.startswith("="):
                cell.data_type = "s"  # text, never a formula
            elif isinstance(value, datetime):
                cell.number_format = "dd mmm yyyy hh:mm"
            elif isinstance(value, date):
                cell.number_format = "dd mmm yyyy"

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = ws.dimensions
    else:
        ws.cell(row=2, column=1, value="Nobody yet.")


async def edition_workbook(cycle: str) -> bytes:
    signed_up = await fetch_all(
        """
        select s.*,
               p.paid_at          as paid_on,
               p.plan_months      as bought_months,
               p.amount_minor     as amount_minor,
               p.currency         as currency,
               p.promo_code       as bought_code,
               p.razorpay_payment_id
        from payments p
        join subscribers s on s.id = p.subscriber_id
        where p.cycle = %s and p.status = 'paid'
        order by p.paid_at
        """,
        (cycle,),
    )
    to_post = await cycles.posting_list(cycle)

    name = cycles.month_name(cycle)
    wb = Workbook()
    _sheet(wb, f"Signed up {name}", SIGNED_UP, signed_up, first=True)
    _sheet(wb, f"To post {name}", TO_POST, to_post, first=False)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
